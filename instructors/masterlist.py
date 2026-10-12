"""Masterlist (CSV / XLSX) parsing, validation and import."""
from __future__ import annotations

import csv
import io
import secrets
from dataclasses import dataclass, field

from django.contrib.auth.hashers import make_password
from django.core.exceptions import ValidationError
from django.db import transaction

from academics.models import BandwidthPolicy, Semester
from accounts.models import User

from .models import SemesterAccount, Instructor

COLUMNS = [
    "username",
    "first_name",
    "middle_name",
    "last_name",
    "extension_name",
]
REQUIRED = ["username", "first_name", "last_name"]
HEADER_ALIASES = {"instructor_number": "username", "instructor_id": "username"}  # older masterlist files
PASSWORD_ALPHABET = "ABCDEFGHJKMNPQRSTUVWXYZabcdefghijkmnpqrstuvwxyz23456789"  # no 0/O/1/l/I


class MasterlistError(Exception):
    """File-level problem (wrong type, missing columns, empty ...)."""


@dataclass
class RowResult:
    line: int
    data: dict
    errors: list[str] = field(default_factory=list)
    status: str = "new"  # new | existing | error

    @property
    def ok(self):
        return not self.errors

    def as_dict(self):
        return {"line": self.line, "data": self.data, "errors": self.errors, "status": self.status}


def generate_password(length=8) -> str:
    return "".join(secrets.choice(PASSWORD_ALPHABET) for _ in range(length))


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _normalise_header(h) -> str:
    name = _cell(h).lower().replace(" ", "_").replace("-", "_")
    return HEADER_ALIASES.get(name, name)


def parse_upload(uploaded) -> list[dict]:
    """Return a list of {column: str} dicts, one per data row, from a CSV or XLSX upload."""
    name = (getattr(uploaded, "name", "") or "").lower()
    raw = uploaded.read()
    if name.endswith(".xlsx"):
        rows = _read_xlsx(raw)
    elif name.endswith(".csv"):
        rows = _read_csv(raw)
    else:
        raise MasterlistError("Unsupported file type. Upload a .csv or .xlsx file (download the template for the format).")
    if not rows:
        raise MasterlistError("The file is empty.")
    header = [_normalise_header(h) for h in rows[0]]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise MasterlistError("Missing required column(s): " + ", ".join(missing) + ". Download the template for the expected format.")
    parsed = []
    for i, row in enumerate(rows[1:], start=2):
        values = {header[j]: _cell(row[j]) if j < len(row) else "" for j in range(len(header))}
        if not any(values.values()):
            continue  # blank line
        parsed.append({"_line": i, **{c: values.get(c, "") for c in COLUMNS}})
    if not parsed:
        raise MasterlistError("The file has a header but no instructor rows.")
    return parsed


def _read_csv(raw: bytes):
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    return list(csv.reader(io.StringIO(text)))


def _read_xlsx(raw: bytes):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - openpyxl raises many types
        raise MasterlistError(f"Could not read the Excel file: {e}") from None
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def validate_rows(parsed: list[dict]) -> list[RowResult]:
    results: list[RowResult] = []
    seen: dict[str, int] = {}
    names = [p["username"] for p in parsed]
    existing = set(Instructor.objects.filter(user__username__in=names).values_list("user__username", flat=True))
    taken_usernames = set(User.objects.filter(username__in=names).values_list("username", flat=True))
    for p in parsed:
        line = p.pop("_line", 0)
        r = RowResult(line=line, data=dict(p))
        for col in REQUIRED:
            if not p[col]:
                r.errors.append(f"Missing {col}")
        username = p["username"]
        if username:
            if len(username) > 150 or " " in username:
                r.errors.append("username must be at most 150 characters with no spaces")
            if username in seen:
                r.errors.append(f"Duplicate username (also on line {seen[username]})")
            else:
                seen[username] = line
        if username in existing:
            r.status = "existing"
        elif username in taken_usernames:
            r.errors.append("Username already used by a non-instructor account")
        if r.errors:
            r.status = "error"
        results.append(r)
    return results


@dataclass
class ImportSummary:
    created: int = 0
    updated: int = 0
    accounts_created: int = 0
    credentials: list[dict] = field(default_factory=list)


def import_rows(valid_rows: list[dict], policy: BandwidthPolicy | None = None) -> ImportSummary:
    """Import validated rows (dicts as produced by RowResult.data)."""
    semester = Semester.get_active()
    if semester is None:
        raise MasterlistError("There is no active semester. Activate a semester before importing a masterlist.")
    policy = policy or BandwidthPolicy.get_default()
    if policy is None:
        raise MasterlistError("There is no bandwidth policy. Create one before importing a masterlist.")

    summary = ImportSummary()
    with transaction.atomic():
        for row in valid_rows:
            username = row["username"]
            instructor = Instructor.objects.select_related("user").filter(user__username=username).first()
            if instructor is None:
                password = generate_password()
                user = User(
                    username=username,
                    first_name=row["first_name"],
                    middle_name=row["middle_name"],
                    last_name=row["last_name"],
                    extension_name=row["extension_name"],
                    is_instructor=True,
                    is_active=True,
                    must_change_password=True,
                )
                user.password = make_password(password, hasher="pbkdf2_sha256_bulk")
                user.save()
                instructor = Instructor.objects.create(user=user)
                summary.created += 1
                summary.credentials.append(
                    {
                        "username": user.username,
                        "name": user.get_full_name(),
                        "password": password,
                    }
                )
            else:
                user = instructor.user
                user.first_name = row["first_name"]
                user.middle_name = row["middle_name"]
                user.last_name = row["last_name"]
                user.extension_name = row["extension_name"]
                user.is_instructor = True
                user.save()
                instructor.approval_status = Instructor.APPROVED  # an administrator listing someone = approved
                instructor.save()
                summary.updated += 1
            _, made = SemesterAccount.objects.get_or_create(instructor=instructor, semester=semester, defaults={"policy": policy})
            summary.accounts_created += int(made)
    return summary


def template_rows() -> list[list[str]]:
    return [COLUMNS, ["jdelacruz", "Juan", "Santos", "Dela Cruz", ""]]
